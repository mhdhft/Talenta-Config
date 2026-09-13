"""
Membuat 2 file Excel dummy (data fiktif, BUKAN data karyawan asli) untuk
testing logic AI mapping Fase 2:

- dummy_data/dummy_new_document.xlsx : kolom acak/berbeda dari standar
- dummy_data/dummy_reference.xlsx    : struktur mengikuti templates/template_A.xlsx

Sengaja dibuat variasi kasus:
- EMP001: Job Position berubah (Changed)
- EMP002: Job Position beda hanya karena spasi tambahan (Need Confirmation)
- EMP003: Job Position typo "Sales Manager" vs "Sales Manger" (Need Confirmation)
- EMP004: Organization Name berubah jelas (Changed)
- EMP005: semua field sama persis (Unchanged)
- EMP006: Branch Name ada simbol aneh "Surabaya#1" (Need Confirmation)
- EMP007: Employment Status End Date berubah (Changed)
- EMP008: Grade berubah (Changed)
- EMP009: Class berubah (Changed)
- EMP010: semua field sama persis (Unchanged)
- EMP011: cuma ada di New Document -> kasus "new_employee" (Not Found)
- EMP012: cuma ada di Reference Document -> kasus "unidentified_employee" (Not Found)
"""

from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = BASE_DIR.parent / "templates"
OUTPUT_DIR = BASE_DIR / "dummy_data"
OUTPUT_DIR.mkdir(exist_ok=True)

NEW_DOCUMENT_HEADERS = [
    "Employee ID",
    "Nama Depan",
    "Nama Belakang",
    "Jabatan",
    "Divisi",
    "Level Jabatan",
    "Grade",
    "Kelas",
    "Sts. Kerja",
    "Tgl Akhir Sts Kerja",
    "Cabang",
    "CC",
    "Kategori CC",
    "Approval Line",
    "Manager",
]

NEW_DOCUMENT_ROWS = [
    ["EMP001", "Andi", "Wijaya", "Senior Software Engineer", "IT Department", 3, "G3", "I", "Permanent", None, "Jakarta", "CC001", "Corporate", "Line A", "Rudi Hartono"],
    ["EMP002", "Siti", "Nurhaliza", "Finance Staff", "Finance", 2, "G2", "II", "Permanent", None, "Bandung", "CC002", "Corporate", "Line B", "Wati Susanti"],
    ["EMP003", "Budi", "Santoso", "Sales Manager", "Sales", 4, "G4", "I", "Permanent", None, "Surabaya", "CC003", "Regional", "Line C", "Doni Prasetyo"],
    ["EMP004", "Rina", "Kartika", "Marketing Staff", "Marketing", 2, "G2", "II", "Permanent", None, "Jakarta", "CC004", "Corporate", "Line A", "Rudi Hartono"],
    ["EMP005", "Dewi", "Anggraini", "HR Staff", "Human Resources", 2, "G2", "II", "Permanent", None, "Jakarta", "CC005", "Corporate", "Line A", "Rudi Hartono"],
    ["EMP006", "Hendra", "Gunawan", "Warehouse Staff", "Logistics", 1, "G1", "III", "Permanent", None, "Surabaya", "CC006", "Regional", "Line C", "Doni Prasetyo"],
    ["EMP007", "Fajar", "Nugroho", "IT Support", "IT Department", 2, "G2", "II", "Contract", "2027-06-30", "Jakarta", "CC001", "Corporate", "Line A", "Rudi Hartono"],
    ["EMP008", "Maya", "Puspita", "Accounting Staff", "Finance", 2, "G4", "II", "Permanent", None, "Bandung", "CC002", "Corporate", "Line B", "Wati Susanti"],
    ["EMP009", "Yusuf", "Ramadhan", "Legal Staff", "Legal", 3, "G3", "II", "Permanent", None, "Jakarta", "CC007", "Corporate", "Line A", "Rudi Hartono"],
    ["EMP010", "Lina", "Marlina", "Customer Service", "Customer Service", 1, "G1", "I", "Permanent", None, "Surabaya", "CC010", "Corporate", "Line C", "Doni Prasetyo"],
    # EMP011 sengaja TIDAK ada di REFERENCE_DATA -> simulasi "karyawan baru".
    ["EMP011", "Bayu", "Saputra", "Marketing Staff", "Marketing", 1, "G1", "I", "Permanent", None, "Jakarta", "CC004", "Corporate", "Line A", "Rudi Hartono"],
]

# Data Reference, key = Employee ID. Field yang tidak disebut otomatis None.
REFERENCE_DATA = {
    "EMP001": {"First Name": "Andi", "Last Name": "Wijaya", "Organization Name": "IT Department", "Job Position": "Software Engineer", "Job Level": 3, "Grade": "G3", "Class": "I", "Employment Status": "Permanent", "Branch Name": "Jakarta", "Cost Center": "CC001", "Cost Center Category": "Corporate", "Approval Line": "Line A", "Manager": "Rudi Hartono"},
    "EMP002": {"First Name": "Siti", "Last Name": "Nurhaliza", "Organization Name": "Finance", "Job Position": "Finance Staff ", "Job Level": 2, "Grade": "G2", "Class": "II", "Employment Status": "Permanent", "Branch Name": "Bandung", "Cost Center": "CC002", "Cost Center Category": "Corporate", "Approval Line": "Line B", "Manager": "Wati Susanti"},
    "EMP003": {"First Name": "Budi", "Last Name": "Santoso", "Organization Name": "Sales", "Job Position": "Sales Manger", "Job Level": 4, "Grade": "G4", "Class": "I", "Employment Status": "Permanent", "Branch Name": "Surabaya", "Cost Center": "CC003", "Cost Center Category": "Regional", "Approval Line": "Line C", "Manager": "Doni Prasetyo"},
    "EMP004": {"First Name": "Rina", "Last Name": "Kartika", "Organization Name": "Marketing & Communication", "Job Position": "Marketing Staff", "Job Level": 2, "Grade": "G2", "Class": "II", "Employment Status": "Permanent", "Branch Name": "Jakarta", "Cost Center": "CC004", "Cost Center Category": "Corporate", "Approval Line": "Line A", "Manager": "Rudi Hartono"},
    "EMP005": {"First Name": "Dewi", "Last Name": "Anggraini", "Organization Name": "Human Resources", "Job Position": "HR Staff", "Job Level": 2, "Grade": "G2", "Class": "II", "Employment Status": "Permanent", "Branch Name": "Jakarta", "Cost Center": "CC005", "Cost Center Category": "Corporate", "Approval Line": "Line A", "Manager": "Rudi Hartono"},
    "EMP006": {"First Name": "Hendra", "Last Name": "Gunawan", "Organization Name": "Logistics", "Job Position": "Warehouse Staff", "Job Level": 1, "Grade": "G1", "Class": "III", "Employment Status": "Permanent", "Branch Name": "Surabaya#1", "Cost Center": "CC006", "Cost Center Category": "Regional", "Approval Line": "Line C", "Manager": "Doni Prasetyo"},
    "EMP007": {"First Name": "Fajar", "Last Name": "Nugroho", "Organization Name": "IT Department", "Job Position": "IT Support", "Job Level": 2, "Grade": "G2", "Class": "II", "Employment Status": "Contract", "Employment Status End Date": "2026-12-31", "Branch Name": "Jakarta", "Cost Center": "CC001", "Cost Center Category": "Corporate", "Approval Line": "Line A", "Manager": "Rudi Hartono"},
    "EMP008": {"First Name": "Maya", "Last Name": "Puspita", "Organization Name": "Finance", "Job Position": "Accounting Staff", "Job Level": 2, "Grade": "G5", "Class": "II", "Employment Status": "Permanent", "Branch Name": "Bandung", "Cost Center": "CC002", "Cost Center Category": "Corporate", "Approval Line": "Line B", "Manager": "Wati Susanti"},
    "EMP009": {"First Name": "Yusuf", "Last Name": "Ramadhan", "Organization Name": "Legal", "Job Position": "Legal Staff", "Job Level": 3, "Grade": "G3", "Class": "III", "Employment Status": "Permanent", "Branch Name": "Jakarta", "Cost Center": "CC007", "Cost Center Category": "Corporate", "Approval Line": "Line A", "Manager": "Rudi Hartono"},
    "EMP010": {"First Name": "Lina", "Last Name": "Marlina", "Organization Name": "Customer Service", "Job Position": "Customer Service", "Job Level": 1, "Grade": "G1", "Class": "I", "Employment Status": "Permanent", "Branch Name": "Surabaya", "Cost Center": "CC010", "Cost Center Category": "Corporate", "Approval Line": "Line C", "Manager": "Doni Prasetyo"},
    # EMP012 sengaja TIDAK ada di NEW_DOCUMENT_ROWS -> simulasi "unidentified employee".
    "EMP012": {"First Name": "Citra", "Last Name": "Dewanti", "Organization Name": "Finance", "Job Position": "Finance Staff", "Job Level": 2, "Grade": "G2", "Class": "II", "Employment Status": "Permanent", "Branch Name": "Bandung", "Cost Center": "CC002", "Cost Center Category": "Corporate", "Approval Line": "Line B", "Manager": "Wati Susanti"},
}


def build_new_document():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(NEW_DOCUMENT_HEADERS)
    for row in NEW_DOCUMENT_ROWS:
        ws.append(row)
    path = OUTPUT_DIR / "dummy_new_document.xlsx"
    wb.save(path)
    return path


def build_reference_document():
    template_wb = openpyxl.load_workbook(TEMPLATES_DIR / "template_A.xlsx")
    headers = [cell.value for cell in template_wb["Worksheet"][1]]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Worksheet"
    ws.append(headers)

    for employee_id, fields in REFERENCE_DATA.items():
        row = [None] * len(headers)
        row[headers.index("Employee ID")] = employee_id
        for field_name, value in fields.items():
            row[headers.index(field_name)] = value
        ws.append(row)

    path = OUTPUT_DIR / "dummy_reference.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    new_path = build_new_document()
    ref_path = build_reference_document()
    print("Dummy files berhasil dibuat:")
    print(" -", new_path)
    print(" -", ref_path)
