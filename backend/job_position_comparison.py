"""
Structure Mapper - Fase D: perbandingan Job Position hasil mapping (Fase B/C,
lihat STRUCTURE-MAPPER-SPEC.md bagian 3.4) terhadap master list Job Position
yang sudah ada di aplikasi Talenta (bagian 3.6 - 3.7, Open Item #2).

Fase E (lihat PROGRESS-MAIN.md) menyambungkan logic di modul ini ke endpoint
backend asli (routers/structure_mapper.py) - fungsi-fungsi di sini DIREUSE
apa adanya, tidak ditulis ulang. `job_position_comparison_test_run.py` (CLI
test manual, Fase D) juga import dari sini, bukan duplikasi logic.

Input:
- mapping_rows: hasil ekstraksi Fase B/C (list Job Position + Parent Job
  Position dari gambar/PDF bagan) - lihat structure_extraction.py.
- master_names: daftar Job Position Name dari sheet "List of Job Position"
  master Talenta - dibaca lewat read_master_job_position_names() di modul
  ini (dipindahkan dari CLI test runner supaya bisa dipakai backend juga).

Output - 3 kategori (plus "sudah ada" sebagai info, lihat spec 3.7):
1. "Baru"                  - ada di mapping, TIDAK ada di master list.
2. "Kandidat Vacant/Dihapus" - ada di master list, TIDAK ada lagi di mapping.
3. "Perlu Konfirmasi"      - nama mirip tapi tidak identik, AI tidak yakin
   sama/beda - JANGAN ditebak (lihat _ai_judge_match).

Algoritma matching (supaya tidak boros panggil AI ke semua kombinasi N x M):
1. Exact match (case-insensitive, whitespace dirapikan) -> langsung "Sudah
   Ada", tidak perlu AI sama sekali.
2. Untuk nama yang tidak exact match: cari 1 kandidat PALING MIRIP di sisa
   master list pakai kemiripan teks (difflib) sebagai penyaring awal. Kalau
   tidak ada kandidat yang cukup mirip (di bawah _FUZZY_CUTOFF) -> langsung
   "Baru", tidak perlu AI (jelas beda).
3. Kalau ada kandidat cukup mirip -> baru di titik ini AI ditanya untuk
   memutuskan "same" (variasi penulisan/typo, sama-sama merujuk posisi yang
   sama) / "different" (memang dua posisi berbeda) / "unsure" (tidak yakin -
   JANGAN ditebak, masuk "Perlu Konfirmasi").
4. Sisa nama master yang di akhir proses tidak terpakai (tidak match ke
   manapun, tidak masuk Perlu Konfirmasi) -> "Kandidat Vacant/Dihapus".

Catatan desain (didiskusikan & disetujui user sebelum coding):
- Kalau 1 nama mapping punya 2+ kandidat mirip di master list, dipakai
  kandidat PALING MIRIP saja (bukan tanya AI untuk semua kandidat) - versi
  pertama, disederhanakan.
- Kolom Status/Job Position Code/Description di master list TIDAK dipakai
  untuk matching - murni berdasarkan Job Position Name (sesuai spec).
"""

import difflib
import io
import re

import openpyxl
from pydantic import BaseModel

from ai_provider import USE_MOCK_AI, call_ai

# Validasi file referensi master list Talenta (sheet "List of Job Position") -
# lihat spec bagian 3.6 & 10.3 (format sheet sama seperti template_excel_for_ai.xlsx).
REQUIRED_MASTER_SHEET_NAME = "List of Job Position"
EXPECTED_MASTER_HEADER = [
    "Job Position Id",
    "Job Position Name",
    "Status",
    "Job Position Code",
    "Description",
]

# Ambang kemiripan teks (0.0 - 1.0, dari difflib.SequenceMatcher.ratio()) untuk
# menentukan apakah 2 nama Job Position "cukup mirip" untuk ditanyakan ke AI.
# Di bawah ambang ini dianggap jelas berbeda, tidak perlu panggil AI (hemat
# kuota) - langsung "Baru" (untuk nama dari mapping) / tetap jadi kandidat
# vacant (untuk nama dari master list).
_FUZZY_CUTOFF = 0.6


def _normalize_header(value) -> str:
    return str(value or "").strip().rstrip("*").strip()


class InvalidMasterListError(ValueError):
    """Dilempar kalau file referensi master list tidak sesuai template (sheet
    tidak ada / header salah) - pesan-nya sudah jelas & siap ditampilkan ke
    user apa adanya (lihat routers/structure_mapper.py & CLI test runner)."""


def read_master_job_position_names(file_bytes: bytes) -> list[str]:
    """Baca daftar Job Position Name dari sheet "List of Job Position" pada
    file Excel referensi (bytes mentah, supaya dipakai baik dari CLI - baca
    dari path - maupun dari endpoint backend - baca dari UploadFile). Validasi
    sheet & header dulu supaya error-nya jelas (bukan gagal diam-diam/crash
    mentah) - lihat spec bagian 3.6, 10.3."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise InvalidMasterListError(
            f"Gagal membaca file Excel master list. Pastikan file tidak rusak/corrupt. ({e})"
        )

    if REQUIRED_MASTER_SHEET_NAME not in wb.sheetnames:
        raise InvalidMasterListError(
            f'Format file tidak sesuai template - sheet "{REQUIRED_MASTER_SHEET_NAME}" tidak ditemukan.'
        )
    ws = wb[REQUIRED_MASTER_SHEET_NAME]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, max_col=5, values_only=True), None)
    normalized_header = [_normalize_header(h) for h in header_row] if header_row else []
    if normalized_header != EXPECTED_MASTER_HEADER:
        raise InvalidMasterListError(
            f'Format file tidak sesuai template - header kolom A-E pada sheet "{REQUIRED_MASTER_SHEET_NAME}" '
            f'harus persis: {", ".join(EXPECTED_MASTER_HEADER)}.'
        )

    names: list[str] = []
    for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        name = row[1]
        if name is None or str(name).strip() == "":
            continue
        names.append(str(name).strip())
    return names


def _normalize_name(value) -> str:
    """Samakan huruf besar/kecil & rapikan whitespace sebelum dibandingkan -
    supaya "Sales Manager" dan " sales   manager " dianggap identik tanpa
    perlu tanya AI sama sekali."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


class JobPositionMatchJudgement(BaseModel):
    verdict: str  # "same" | "different" | "unsure"
    alasan: str


def _mock_judge_match(name_mapping: str, name_master: str) -> JobPositionMatchJudgement:
    """Mode testing tanpa panggil Gemini (lihat USE_MOCK_AI di ai_provider.py).
    Pakai kemiripan teks sederhana sebagai pengganti judgment AI asli - makin
    mirip makin condong "same", makin beda makin condong "different", di
    tengah-tengah "unsure"."""
    ratio = difflib.SequenceMatcher(None, name_mapping.lower(), name_master.lower()).ratio()
    if ratio >= 0.82:
        verdict = "same"
    elif ratio <= 0.5:
        verdict = "different"
    else:
        verdict = "unsure"
    return JobPositionMatchJudgement(
        verdict=verdict,
        alasan=f"[MOCK] Kemiripan teks {ratio:.0%} antara kedua nama (USE_MOCK_AI=true, tidak memanggil Gemini).",
    )


def _ai_judge_match(name_mapping: str, name_master: str) -> JobPositionMatchJudgement:
    if USE_MOCK_AI:
        return _mock_judge_match(name_mapping, name_master)

    prompt = f"""Kamu membandingkan dua nama Job Position (jabatan) dari dua sumber data
HR yang berbeda, untuk menentukan apakah keduanya MERUJUK POSISI YANG SAMA
(cuma beda penulisan/singkatan/urutan kata/typo) atau memang DUA POSISI YANG
BERBEDA (beda level/fungsi/departemen).

Nama Job Position 1 (hasil pemetaan bagan struktur terbaru): "{name_mapping}"
Nama Job Position 2 (master list existing di aplikasi Talenta): "{name_master}"

Tentukan salah satu verdict:
- "same" - yakin ini posisi yang sama, cuma variasi penulisan/singkatan/typo.
- "different" - yakin ini dua posisi yang berbeda secara makna/level/fungsi.
- "unsure" - ragu, tidak cukup yakin untuk memutuskan sama atau beda. JANGAN
  menebak kalau ragu - pilih ini kalau memang tidak yakin.

Kembalikan verdict beserta alasan singkat (field "alasan")."""

    return call_ai(prompt, response_schema=JobPositionMatchJudgement)


def compare_job_positions(mapping_rows: list[dict], master_names: list[str]) -> dict:
    """
    mapping_rows: [{"job_position": str, "parent_job_position": str | None}, ...]
        - hasil ekstraksi Fase B/C, SEMUA baris apa adanya seperti tampil di
          laporan mapping 3.4 (termasuk baris yang masih ditandai "perlu
          dicek manual" dari Fase B - nama posisinya tetap ikut dibandingkan).
    master_names: [str, ...] - daftar Job Position Name dari sheet "List of
        Job Position" master Talenta (3.6). Boleh mengandung nama duplikat.

    Return: {
        "sudah_ada": [{"job_position": str, "parent_job_position": str|None}, ...],
        "baru": [{"job_position": str, "parent_job_position": str|None}, ...],
        "vacant_candidates": [{"job_position": str}, ...],
        "perlu_konfirmasi": [
            {"job_position_mapping": str, "parent_job_position": str|None,
             "job_position_master": str, "alasan": str},
            ...
        ],
    }
    """
    sudah_ada: list[dict] = []
    baru: list[dict] = []
    perlu_konfirmasi: list[dict] = []

    # Pool sisa master list yang belum "terpakai" (belum jadi Sudah Ada /
    # Perlu Konfirmasi) - list of dict {"name": original, "norm": normalized},
    # supaya bisa dihapus satu-per-satu (konsumsi) selagi diproses, dan sisa
    # di akhir otomatis jadi "Kandidat Vacant/Dihapus".
    remaining_master = [{"name": name, "norm": _normalize_name(name)} for name in master_names]

    for item in mapping_rows:
        job_position = item.get("job_position")
        parent_job_position = item.get("parent_job_position")
        norm_mapping = _normalize_name(job_position)

        # Langkah 1: exact match (case-insensitive, whitespace dirapikan).
        exact_index = next(
            (i for i, m in enumerate(remaining_master) if m["norm"] == norm_mapping), None
        )
        if exact_index is not None:
            matched = remaining_master.pop(exact_index)
            sudah_ada.append(
                {"job_position": job_position, "parent_job_position": parent_job_position}
            )
            continue

        # Langkah 2: cari 1 kandidat paling mirip di sisa master list.
        if remaining_master:
            candidate_norms = [m["norm"] for m in remaining_master]
            close = difflib.get_close_matches(norm_mapping, candidate_norms, n=1, cutoff=_FUZZY_CUTOFF)
        else:
            close = []

        if not close:
            # Tidak ada kandidat yang cukup mirip - jelas beda, tidak perlu AI.
            baru.append(
                {"job_position": job_position, "parent_job_position": parent_job_position}
            )
            continue

        candidate_index = next(
            i for i, m in enumerate(remaining_master) if m["norm"] == close[0]
        )
        candidate_name = remaining_master[candidate_index]["name"]

        # Langkah 3: kandidat cukup mirip tapi tidak identik - baru di titik
        # ini AI ditanya (hemat kuota - tidak setiap pasangan ditanyakan).
        judgement = _ai_judge_match(job_position, candidate_name)

        if judgement.verdict == "same":
            remaining_master.pop(candidate_index)
            sudah_ada.append(
                {"job_position": job_position, "parent_job_position": parent_job_position}
            )
        elif judgement.verdict == "different":
            # Kandidat TIDAK dikonsumsi - tetap ada di pool, bisa saja masih
            # match ke item mapping lain, atau berakhir jadi vacant candidate.
            baru.append(
                {"job_position": job_position, "parent_job_position": parent_job_position}
            )
        else:  # "unsure" - jangan ditebak, masuk Perlu Konfirmasi.
            remaining_master.pop(candidate_index)
            perlu_konfirmasi.append(
                {
                    "job_position_mapping": job_position,
                    "parent_job_position": parent_job_position,
                    "job_position_master": candidate_name,
                    "alasan": judgement.alasan,
                }
            )

    vacant_candidates = [{"job_position": m["name"]} for m in remaining_master]

    return {
        "sudah_ada": sudah_ada,
        "baru": baru,
        "vacant_candidates": vacant_candidates,
        "perlu_konfirmasi": perlu_konfirmasi,
    }


def build_new_positions_export_rows(baru_rows: list[dict]) -> list[dict]:
    """Siapkan data untuk export Excel kategori "Baru" (spec 3.7 poin 1) -
    2 kolom saja: Job Position Name & Parent Job Position Name, siap
    diimport ke Talenta. Diambil langsung dari data hasil mapping (bukan
    dari master list)."""
    return [
        {
            "Job Position Name": row["job_position"],
            "Parent Job Position Name": row["parent_job_position"],
        }
        for row in baru_rows
    ]
