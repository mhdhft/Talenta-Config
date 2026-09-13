"""
Structure Mapper - Fase C: integrasi backend untuk step 3.2-3.4 (upload ->
AI analysis -> report mapping). Lihat STRUCTURE-MAPPER-SPEC.md bagian 5, 6
(Open Item #1) dan 3.8.

Fase E (lihat PROGRESS-MAIN.md) menambahkan step 3.6-3.7 (upload Referensi
Talenta -> bandingkan -> report perbandingan 3 kategori) di file yang sama
ini - reuse logic Fase D (job_position_comparison.py) apa adanya, tidak
ditulis ulang.

Reuse logic Fase B (structure_extraction.py) apa adanya - tidak ditulis
ulang. Validasi format file (PDF/JPG/JPEG/PNG) sengaja TERPISAH dari
validator .xlsx/.csv di routers/documents.py (lihat guess_mime_type di
structure_extraction.py, dipakai juga untuk validasi ini).
"""

import io
import re
import shutil
import uuid
from pathlib import Path
from typing import Literal

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from deps import get_current_user, get_db
from job_position_comparison import (
    InvalidMasterListError,
    build_new_positions_export_rows,
    compare_job_positions,
    read_master_job_position_names,
)
from models import (
    JobPositionComparisonRow,
    JobPositionComparisonSession,
    StructureMapperRow,
    StructureMapperSession,
    User,
)
from schemas import (
    ExportComparisonListRequest,
    JobPositionComparisonResponse,
    JobPositionComparisonRowOut,
    StructureMapperRowOut,
    StructureMapperUploadResponse,
)
from structure_extraction import BoxStatus, StructureRow, extract_structure, guess_mime_type

router = APIRouter(prefix="/structure-mapper", tags=["structure-mapper"])

BASE_DIR = Path(__file__).resolve().parent.parent
UPLOAD_DIR = BASE_DIR / "uploads" / "structure-mapper"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

# Terjemahan status internal Fase B (structure_extraction.BoxStatus, 4 nilai)
# ke vocabulary status yang lebih sederhana untuk DB/API/UI (3 nilai) - lihat
# PROGRESS-MAIN.md Fase C untuk alasan penyederhanaan ini.
_STATUS_MAP: dict[BoxStatus, str] = {
    "ok": "ok",
    "ambiguous_solid_merge": "duplicated",
    "dotted_line_multi_parent": "needs_manual",
    "failed_unreadable": "needs_manual",
}


def _save_upload(file: UploadFile, ext: str) -> Path:
    dest = UPLOAD_DIR / f"{uuid.uuid4().hex}{ext}"
    with dest.open("wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return dest


def _row_to_db_status(row: StructureRow) -> str:
    return _STATUS_MAP.get(row.status, "needs_manual")


def _get_session(db: Session, session_id: int, user: User) -> StructureMapperSession:
    session = db.get(StructureMapperSession, session_id)
    if session is None or session.user_id != user.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sesi Structure Mapper tidak ditemukan.")
    return session


def _get_comparison_session(
    db: Session, comparison_session_id: int, user: User
) -> JobPositionComparisonSession:
    """Sama seperti _get_session, tapi untuk sesi perbandingan Fase E -
    kepemilikan dicek lewat StructureMapperSession yang terkait (comparison
    session sendiri tidak punya user_id langsung, lihat models.py)."""
    comparison_session = db.get(JobPositionComparisonSession, comparison_session_id)
    if comparison_session is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sesi perbandingan tidak ditemukan.")
    structure_session = db.get(
        StructureMapperSession, comparison_session.structure_mapper_session_id
    )
    if structure_session is None or structure_session.user_id != user.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sesi perbandingan tidak ditemukan.")
    return comparison_session


@router.post("/job-position/upload", response_model=StructureMapperUploadResponse)
def upload_job_position(
    cid: str = Form(...),
    company_name: str = Form(...),
    file: UploadFile = File(...),
    ignore_unsolid_line: bool = Form(False),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    cid = cid.strip()
    company_name = company_name.strip()
    if not cid or not company_name:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Company ID dan Company Name wajib diisi.")

    try:
        mime_type = guess_mime_type(file.filename or "")
    except ValueError as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(e))

    ext = Path(file.filename or "").suffix.lower()
    dest = _save_upload(file, ext)
    file_bytes = dest.read_bytes()

    # include_names selalu True untuk sekarang - belum ada UI untuk user
    # mematikan pembacaan nama (lihat PROGRESS-MAIN.md Fase C, keputusan #1).
    # Toggle "Tampilkan kolom Name" di frontend murni show/hide tampilan atas
    # data yang sudah kembali dari sini.
    include_names = True

    try:
        rows = extract_structure(
            file_bytes,
            mime_type,
            include_names=include_names,
            source_filename=file.filename,
            ignore_unsolid_line=ignore_unsolid_line,
        )
    except Exception as e:
        dest.unlink(missing_ok=True)
        raise HTTPException(
            status.HTTP_502_BAD_GATEWAY,
            f"Gagal menganalisa dokumen dengan AI. Silakan coba lagi. ({e})",
        )

    # Tiap submit = SESI BARU (histori tetap ada di DB, tidak overwrite -
    # lihat spec bagian 3.8 & PROGRESS-MAIN.md Fase C keputusan reset).
    session = StructureMapperSession(
        user_id=current_user.id,
        type="job_position",
        cid=cid,
        company_name=company_name,
        source_filename=file.filename or dest.name,
        source_file_path=str(dest),
        include_names=include_names,
        ignore_unsolid_line=ignore_unsolid_line,
    )
    db.add(session)
    db.flush()  # supaya session.id terisi sebelum dipakai di row di bawah

    row_models: list[StructureMapperRow] = []
    for row in rows:
        row_model = StructureMapperRow(
            session_id=session.id,
            job_position=row.job_title,
            parent_job_position=row.parent,
            name=row.name,
            status=_row_to_db_status(row),
            note=row.note,
            source_box_id_internal=row.source_box_id_internal,
        )
        db.add(row_model)
        row_models.append(row_model)

    db.commit()
    for row_model in row_models:
        db.refresh(row_model)

    return StructureMapperUploadResponse(
        session_id=session.id,
        cid=session.cid,
        company_name=session.company_name,
        source_filename=session.source_filename,
        rows=[
            StructureMapperRowOut(
                id=r.id,
                job_position=r.job_position,
                parent_job_position=r.parent_job_position,
                name=r.name,
                status=r.status,
                note=r.note,
            )
            for r in row_models
        ],
    )


@router.get("/job-position/{session_id}/export")
def export_job_position(
    session_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate Excel on-the-fly dari hasil mapping sesi ini (belum pakai
    styling/Template B - itu Fase E, lihat PROGRESS-MAIN.md Fase C
    keputusan #2). Tidak disimpan permanen di server (beda dari
    routers/report.py yang punya retensi 2 file terbaru - aturan itu
    khusus flow utama, tidak dipakai di sini)."""
    session = _get_session(db, session_id, current_user)
    rows = (
        db.query(StructureMapperRow)
        .filter(StructureMapperRow.session_id == session.id)
        .order_by(StructureMapperRow.id)
        .all()
    )
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Belum ada hasil mapping untuk sesi ini.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job Position Mapping"
    ws.append(["Job Position", "Parent Job Position", "Name", "Status", "Note"])
    for row in rows:
        ws.append([row.job_position, row.parent_job_position, row.name, row.status, row.note])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    safe_company = re.sub(r'[/\\:*?"<>|\s]+', "_", session.company_name.strip())
    safe_cid = re.sub(r'[/\\:*?"<>|\s]+', "_", session.cid.strip())
    filename = f"StructureMapper_{safe_cid}-{safe_company}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---- Fase E: step 3.6-3.7 - upload Referensi Talenta -> bandingkan -> report
# 3 kategori. Reuse job_position_comparison.py (Fase D) apa adanya. ----


@router.post("/job-position/{session_id}/compare", response_model=JobPositionComparisonResponse)
def compare_job_position(
    session_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Step 3.6-3.7 (lihat STRUCTURE-MAPPER-SPEC.md): terima file referensi
    master list Talenta, bandingkan dengan SEMUA baris hasil mapping sesi ini
    (apa adanya, termasuk yang masih "perlu dicek manual" dari Fase B - lihat
    keputusan Fase D), pakai compare_job_positions() (Fase D, TIDAK ditulis
    ulang). Tiap submit = sesi perbandingan BARU (histori tetap ada di DB,
    bukan overwrite - sama seperti pola StructureMapperSession, penting juga
    untuk test konsistensi: jalankan proses yang sama berkali-kali harus tetap
    bisa dibandingkan hasilnya per-run)."""
    session = _get_session(db, session_id, current_user)

    mapping_rows_db = (
        db.query(StructureMapperRow)
        .filter(StructureMapperRow.session_id == session.id)
        .order_by(StructureMapperRow.id)
        .all()
    )
    if not mapping_rows_db:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "Belum ada hasil mapping untuk sesi ini."
        )
    mapping_rows = [
        {"job_position": r.job_position, "parent_job_position": r.parent_job_position}
        for r in mapping_rows_db
    ]

    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx", ".xls"):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Format file tidak didukung. Upload file Excel (.xlsx atau .xls).",
        )

    dest = _save_upload(file, ext)
    file_bytes = dest.read_bytes()

    try:
        master_names = read_master_job_position_names(file_bytes)
    except InvalidMasterListError as e:
        dest.unlink(missing_ok=True)
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(e))

    try:
        result = compare_job_positions(mapping_rows, master_names)
    except Exception as e:
        dest.unlink(missing_ok=True)
        raise HTTPException(
            status.HTTP_502_BAD_GATEWAY,
            f"Gagal memproses perbandingan dengan AI. Silakan coba lagi. ({e})",
        )

    comparison_session = JobPositionComparisonSession(
        structure_mapper_session_id=session.id,
        reference_filename=file.filename or dest.name,
        reference_file_path=str(dest),
    )
    db.add(comparison_session)
    db.flush()  # supaya comparison_session.id terisi sebelum dipakai di row di bawah

    row_models: list[JobPositionComparisonRow] = []
    for item in result["sudah_ada"]:
        row_models.append(
            JobPositionComparisonRow(
                comparison_session_id=comparison_session.id,
                category="sudah_ada",
                job_position=item["job_position"],
                parent_job_position=item["parent_job_position"],
            )
        )
    for item in result["baru"]:
        row_models.append(
            JobPositionComparisonRow(
                comparison_session_id=comparison_session.id,
                category="baru",
                job_position=item["job_position"],
                parent_job_position=item["parent_job_position"],
            )
        )
    for item in result["vacant_candidates"]:
        row_models.append(
            JobPositionComparisonRow(
                comparison_session_id=comparison_session.id,
                category="vacant_candidate",
                job_position=item["job_position"],
            )
        )
    for item in result["perlu_konfirmasi"]:
        row_models.append(
            JobPositionComparisonRow(
                comparison_session_id=comparison_session.id,
                category="perlu_konfirmasi",
                job_position=item["job_position_mapping"],
                parent_job_position=item["parent_job_position"],
                master_job_position=item["job_position_master"],
                alasan=item["alasan"],
            )
        )
    for row_model in row_models:
        db.add(row_model)

    db.commit()
    for row_model in row_models:
        db.refresh(row_model)

    return JobPositionComparisonResponse(
        comparison_session_id=comparison_session.id,
        structure_mapper_session_id=session.id,
        reference_filename=comparison_session.reference_filename,
        rows=[
            JobPositionComparisonRowOut(
                id=r.id,
                category=r.category,
                job_position=r.job_position,
                parent_job_position=r.parent_job_position,
                master_job_position=r.master_job_position,
                alasan=r.alasan,
            )
            for r in row_models
        ],
    )


@router.get("/job-position/comparison/{comparison_session_id}/export-new")
def export_new_positions(
    comparison_session_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export Excel khusus kategori "Baru" (spec 3.7 poin 1) - 2 kolom saja:
    Job Position Name & Parent Job Position Name, siap diimport ke Talenta.
    Reuse build_new_positions_export_rows() (Fase D) apa adanya."""
    comparison_session = _get_comparison_session(db, comparison_session_id, current_user)
    baru_rows_db = (
        db.query(JobPositionComparisonRow)
        .filter(
            JobPositionComparisonRow.comparison_session_id == comparison_session.id,
            JobPositionComparisonRow.category == "baru",
        )
        .order_by(JobPositionComparisonRow.id)
        .all()
    )
    if not baru_rows_db:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "Tidak ada Job Position kategori 'Baru' untuk sesi ini."
        )

    export_rows = build_new_positions_export_rows(
        [
            {"job_position": r.job_position, "parent_job_position": r.parent_job_position}
            for r in baru_rows_db
        ]
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "New Job Positions"
    ws.append(["Job Position Name", "Parent Job Position Name"])
    for row in export_rows:
        ws.append([row["Job Position Name"], row["Parent Job Position Name"]])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    structure_session = db.get(
        StructureMapperSession, comparison_session.structure_mapper_session_id
    )
    safe_company = re.sub(r'[/\\:*?"<>|\s]+', "_", structure_session.company_name.strip())
    safe_cid = re.sub(r'[/\\:*?"<>|\s]+', "_", structure_session.cid.strip())
    filename = f"StructureMapper_NewPositions_{safe_cid}-{safe_company}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---- Revisi 2026-08-11 (spec 3.7b): Step 5 - report "Baru & Sudah Ada" (di
# atas) dipecah dari "Vacant & Perlu Konfirmasi" (di bawah). Keputusan radio
# per baris (Tetap Vacant/Set Inactive/Anggap Sama) sengaja disimpan di
# FRONTEND SAJA (state React, bukan DB - sama seperti pola checklist Fase E
# sebelumnya), jadi endpoint ini cuma menerima daftar nama yang dipilih user
# saat tombol download diklik, lalu generate Excel-nya on the fly. ----


@router.post("/job-position/comparison/{comparison_session_id}/export-list/{kind}")
def export_comparison_list(
    comparison_session_id: int,
    kind: Literal["vacant", "inactive"],
    payload: ExportComparisonListRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export "Download List Vacant" / "Download List Inactive" (spec 3.7b).
    Format: 1 kolom saja (Job Position Name) - Id/Status/Parent belum
    diperlukan, keputusan format detail menyusul kalau dibutuhkan.

    Nama yang dikirim dari browser disaring dulu terhadap baris kategori
    "vacant_candidate"/"perlu_konfirmasi" yang beneran ada di sesi
    perbandingan ini - supaya tidak asal percaya input mentah dari client
    (misal sesi lama/basi)."""
    comparison_session = _get_comparison_session(db, comparison_session_id, current_user)

    valid_rows = (
        db.query(JobPositionComparisonRow)
        .filter(
            JobPositionComparisonRow.comparison_session_id == comparison_session.id,
            JobPositionComparisonRow.category.in_(["vacant_candidate", "perlu_konfirmasi"]),
        )
        .all()
    )
    valid_names = {r.job_position for r in valid_rows}
    # Pertahankan urutan yang dikirim browser, saring nama yang tidak valid.
    selected_names = [name for name in payload.job_positions if name in valid_names]
    if not selected_names:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Tidak ada Job Position yang dipilih untuk kategori ini.",
        )

    label = "Vacant" if kind == "vacant" else "Inactive"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"List {label}"
    ws.append(["Job Position Name"])
    for name in selected_names:
        ws.append([name])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    structure_session = db.get(
        StructureMapperSession, comparison_session.structure_mapper_session_id
    )
    safe_company = re.sub(r'[/\\:*?"<>|\s]+', "_", structure_session.company_name.strip())
    safe_cid = re.sub(r'[/\\:*?"<>|\s]+', "_", structure_session.cid.strip())
    filename = f"StructureMapper_{label}_{safe_cid}-{safe_company}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
