"""
Structure Mapper - "Structure Your Excel" (dulu "Organization") - Fase F2:
parsing file Excel asli (format tetap/fixed) jadi diagram struktur. Lihat
STRUCTURE-MAPPER-SPEC.md bagian 10 untuk spec lengkap.

Router INI SENGAJA TERPISAH dari routers/structure_mapper.py (yang menangani
flow "Excel Your Structure" / dulu "Job Position") - dua flow ini independen,
beda arah transformasi data, tidak saling bergantung (lihat spec bagian 10.1
& Section 7 Dependency Notes). Jangan gabungkan/edit file itu untuk fitur ini.

Beda dari flow "Excel Your Structure": fitur ini TIDAK PAKAI AI SAMA SEKALI -
format Excel-nya selalu sama persis (fixed template), jadi parsing-nya murni
deterministik (baca kolom, bangun tree) pakai openpyxl saja.

Fase F2 ini juga SENGAJA TIDAK menyimpan apapun ke database (session-only) -
keputusan eksplisit user, lihat spec bagian 10.5. Kalau nanti histori/
penyimpanan permanen dibutuhkan, itu perubahan terpisah yang perlu didiskusikan
dulu.
"""

import io
from pathlib import Path

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status

from deps import get_current_user
from models import User
from schemas import OrganizationUploadResponse, OrgTreeNodeOut

router = APIRouter(prefix="/structure-mapper/organization", tags=["structure-mapper-organization"])

REQUIRED_SHEET_NAME = "Export or Import File"

# Header kolom A-D (lihat spec bagian 10.3). Dibandingkan setelah di-strip
# tanda "*" di akhir dan whitespace - file contoh asli (template_excel_for_ai.xlsx)
# menulis "Job Position Id*"/"Job Position Name*" (tanda bintang = kolom
# wajib), sementara kolom C/D tidak pakai tanda itu. Validasi header di sini
# sengaja toleran terhadap ada/tidaknya "*" itu, supaya tidak rapuh cuma
# gara-gara perbedaan kecil semacam itu - fokus validasinya ke NAMA kolom.
EXPECTED_HEADERS = [
    "Job Position Id",
    "Job Position Name",
    "Parent Job Position Id",
    "Parent Job Position Name",
]


def _normalize_header(value: object) -> str:
    return str(value or "").strip().rstrip("*").strip()


def _normalize_id(value: object) -> str | None:
    """Job Position Id/Parent Id dibaca sebagai teks internal untuk matching -
    tidak pernah ditampilkan ke user (lihat spec bagian 10.3). Kalau openpyxl
    kebetulan membaca cell sebagai angka (float), buang ".0"-nya supaya
    "1164325.0" tidak dianggap beda dari "1164325"."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _normalize_name(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


@router.post("/upload", response_model=OrganizationUploadResponse)
def upload_organization(
    cid: str = Form(...),
    company_name: str = Form(...),
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    cid = cid.strip()
    company_name = company_name.strip()
    if not cid or not company_name:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Company ID dan Company Name wajib diisi.")

    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx", ".xls"):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Format file tidak didukung. Upload file Excel (.xlsx atau .xls).",
        )

    file_bytes = file.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Gagal membaca file Excel. Pastikan file tidak rusak/corrupt. ({e})",
        )

    if REQUIRED_SHEET_NAME not in wb.sheetnames:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f'Format file tidak sesuai template - sheet "{REQUIRED_SHEET_NAME}" tidak ditemukan.',
        )
    ws = wb[REQUIRED_SHEET_NAME]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, max_col=4, values_only=True), None)
    normalized_header = [_normalize_header(h) for h in header_row] if header_row else []
    if normalized_header != EXPECTED_HEADERS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Format file tidak sesuai template - header kolom A-D pada sheet "
            f'"{REQUIRED_SHEET_NAME}" harus persis: {", ".join(EXPECTED_HEADERS)}.',
        )

    # Baca semua baris data. Kolom E (Job Position Code) & F (Description)
    # SENGAJA diabaikan total - tidak dibaca sama sekali (lihat spec 10.3).
    records: list[dict] = []
    for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        pos_id_raw, pos_name_raw, parent_id_raw, parent_name_raw = row
        pos_id = _normalize_id(pos_id_raw)
        pos_name = _normalize_name(pos_name_raw)
        parent_id = _normalize_id(parent_id_raw)
        if pos_id is None and pos_name is None:
            continue  # baris kosong sepenuhnya - lewati, bukan data
        records.append(
            {
                "id": pos_id,
                "name": pos_name or "(tanpa nama)",
                "parent_id": parent_id,
            }
        )

    if not records:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "File tidak berisi data job position apapun."
        )

    ids = {r["id"] for r in records if r["id"] is not None}
    children_by_parent: dict[str, list[dict]] = {}
    warnings: list[str] = []
    roots: list[dict] = []

    for r in records:
        parent_id = r["parent_id"]
        if not parent_id:
            roots.append(r)
        elif parent_id not in ids:
            # Referensi rusak (Parent Job Position Id tidak ditemukan) -
            # JANGAN crash, tampilkan sebagai warning yang sebut NAMA baris
            # (bukan Id, karena Id tidak pernah ditampilkan ke user) dan
            # perlakukan baris ini sebagai root terpisah supaya datanya tetap
            # kelihatan (bukan hilang diam-diam) - lihat spec bagian 10.3.
            warnings.append(
                f'Posisi "{r["name"]}" mereferensikan atasan yang tidak ditemukan di '
                "file (referensi rusak) - ditampilkan sebagai pohon/root terpisah."
            )
            roots.append(r)
        else:
            children_by_parent.setdefault(parent_id, []).append(r)

    # "visited" GLOBAL (bukan cuma per-path) - dipakai untuk mendeteksi grup
    # referensi melingkar yang SAMA SEKALI TIDAK TERHUBUNG ke root manapun
    # (mis. A<->B saling menunjuk satu sama lain, tapi tidak ada yang jadi
    # root/anak dari root manapun) - tanpa ini, grup semacam itu tidak akan
    # pernah "dikunjungi" mulai dari `roots`, sehingga hilang dari hasil
    # TANPA warning sama sekali (persis silent-fail yang mau dihindari).
    visited: set[str] = set()

    def build_node(record: dict, path_seen: frozenset[str]) -> OrgTreeNodeOut:
        node_id = record["id"] or f"_noid_{id(record)}"
        visited.add(node_id)
        if node_id in path_seen:
            # Proteksi tambahan (di luar yang diminta eksplisit di spec) -
            # cegah infinite loop kalau data punya referensi melingkar
            # (mis. A punya parent B, B punya parent A). Dihentikan di sini,
            # bukan crash, dan user diberi tahu lewat warning.
            warnings.append(
                f'Terdeteksi referensi melingkar (circular reference) pada posisi "{record["name"]}" '
                "- cabang ini dihentikan di titik tsb untuk mencegah loop tak berujung."
            )
            return OrgTreeNodeOut(id=node_id, name=record["name"], children=[])
        next_path_seen = path_seen | {node_id}
        children = (
            [build_node(c, next_path_seen) for c in children_by_parent.get(record["id"], [])]
            if record["id"]
            else []
        )
        return OrgTreeNodeOut(id=node_id, name=record["name"], children=children)

    trees = [build_node(r, frozenset()) for r in roots]

    # Sisa record yang TIDAK PERNAH terkunjungi lewat root manapun di atas -
    # ini grup referensi melingkar yang terisolasi total (lihat komentar
    # "visited" di atas). Jadikan masing-masing root tambahan (supaya datanya
    # tetap tampil, bukan hilang), tandai jelas ke user lewat warning.
    for r in records:
        node_id = r["id"] or f"_noid_{id(r)}"
        if node_id in visited:
            continue
        warnings.append(
            f'Terdeteksi grup referensi melingkar (circular reference) yang tidak terhubung ke '
            f'root manapun, dimulai dari posisi "{r["name"]}" - ditampilkan sebagai pohon/root terpisah.'
        )
        trees.append(build_node(r, frozenset()))

    return OrganizationUploadResponse(
        cid=cid,
        company_name=company_name,
        source_filename=file.filename or "structure.xlsx",
        trees=trees,
        warnings=warnings,
    )
